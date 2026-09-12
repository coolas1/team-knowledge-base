"""Upstream image-slide assembly plus independent office rendering validation."""

import hashlib
import importlib.util
import json
from pathlib import Path
import shutil
import subprocess
import tempfile

from pptx import Presentation
from pypdf import PdfReader

from .provider import validate_image
from .verify_vendor import verify


def validate_pptx(path: Path, pages: list, spec: dict):
    deck = Presentation(path)
    if len(deck.slides) != len(pages) or deck.slide_width * 9 != deck.slide_height * 16:
        raise ValueError("PPTX page count or aspect ratio mismatch")
    for n, slide in enumerate(deck.slides):
        if len(slide.shapes) != 1 or not hasattr(slide.shapes[0], "image"):
            raise ValueError("PPTX image missing")
        shape = slide.shapes[0]
        if (shape.left, shape.top, shape.width, shape.height) != (
            0,
            0,
            deck.slide_width,
            deck.slide_height,
        ):
            raise ValueError("PPTX image crop or placement mismatch")
        if slide.notes_slide.notes_text_frame.text != spec["pages"][n]["notes"]:
            raise ValueError("PPTX speaker notes mismatch")


def render_pptx(path: Path, expected_pages: int):
    output = path.parent / "rendered"
    output.mkdir(exist_ok=True)
    office = shutil.which("libreoffice") or shutil.which("soffice")
    rasterizer = shutil.which("pdftoppm")
    if not office or not rasterizer:
        raise RuntimeError("Office renderer is unavailable")
    with tempfile.TemporaryDirectory(prefix="tkb-ppt-office-") as profile:
        subprocess.run(
            [
                office,
                "-env:UserInstallation=" + Path(profile).as_uri(),
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output),
                str(path),
            ],
            check=True,
            timeout=90,
            capture_output=True,
        )
    pdf = output / (path.stem + ".pdf")
    if not pdf.is_file() or len(PdfReader(pdf).pages) != expected_pages:
        raise ValueError("Office render page count mismatch")
    subprocess.run(
        [rasterizer, "-scale-to", "1600", "-png", str(pdf), str(output / "page")],
        check=True,
        timeout=90,
        capture_output=True,
    )
    renders = list(output.glob("page-*.png"))
    if len(renders) != expected_pages:
        raise ValueError("Office preview pages missing")
    for image in renders:
        validate_image(image.read_bytes())
    return {
        "tool": subprocess.check_output([office, "--version"], text=True).strip(),
        "pages": expected_pages,
    }


class Assembler:
    def __init__(self, store):
        self.store = store

    def __call__(self, job, pages):
        verify()
        if len(pages) != len(job.spec["pages"]) or not all(
            p.status == "accepted" and p.qa and p.qa.get("passed") is True
            for p in pages
        ):
            raise ValueError("Every page must pass visual QA")
        folder = self.store.path(f"{job.id}/revision-{job.revision}")
        originals = folder / "origin_image"
        originals.mkdir(parents=True, exist_ok=True)
        images = []
        for p in pages:
            source = self.store.path(p.result["path"])
            data = source.read_bytes()
            if hashlib.sha256(data).hexdigest() != p.result["sha256"]:
                raise ValueError("Source image checksum mismatch")
            mime = validate_image(data, slide=True)
            destination = (
                originals
                / f"slide_{p.number:02d}.{'png' if mime == 'image/png' else 'jpg'}"
            )
            destination.write_bytes(data)
            images.append(str(destination))
        (folder / "outline.md").write_text(
            "\n\n".join(
                f"## {i + 1}. {p['title']}\n" + "\n".join("- " + x for x in p["points"])
                for i, p in enumerate(job.spec["pages"])
            ),
            encoding="utf-8",
        )
        (folder / "speech.md").write_text(
            "\n\n".join(
                f"## Slide {i + 1}: {p['title']}\n\n{p['notes']}"
                for i, p in enumerate(job.spec["pages"])
            ),
            encoding="utf-8",
        )
        for name, value in {
            "deck_spec.json": job.spec,
            "slide_jobs.json": [
                {
                    "number": p.number,
                    "status": "accepted",
                    "result": p.result,
                    "qa": p.qa,
                }
                for p in pages
            ],
        }.items():
            (folder / name).write_text(
                json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8"
            )
        module_path = Path(__file__).parent / "vendor/codex_ppt/scripts/assemble_ppt.py"
        loader = importlib.util.spec_from_file_location(
            "tkb_pinned_ppt_assembly", module_path
        )
        module = importlib.util.module_from_spec(loader)
        loader.loader.exec_module(module)
        path = folder / "presentation.pptx"
        if not module.create_presentation(
            images,
            str(path),
            "16:9",
            {i + 1: p["notes"] for i, p in enumerate(job.spec["pages"])},
        ):
            raise ValueError("Upstream PPT assembly failed")
        validate_pptx(path, pages, job.spec)
        rendered = render_pptx(path, len(pages))
        (folder / "validation.json").write_text(json.dumps(rendered), encoding="utf-8")
        return {
            "path": str(path.relative_to(self.store.root)).replace("\\", "/"),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "pages": len(pages),
            "download_url": f"/api/ppt/jobs/{job.id}/download",
            "render": rendered,
        }
